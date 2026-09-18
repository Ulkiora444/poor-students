import numpy as np
import os, os.path
from fastapi import FastAPI, Body
from fastapi.middleware.cors import CORSMiddleware
import time
import requests
from docxtpl import DocxTemplate
import win32api, win32print, os, tempfile
import docx
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Pt, Inches
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
import requests
from docx.oxml.ns import qn 
from docx.oxml import OxmlElement

def rotate_text(cell): 
        tc = cell._element 
        tcPr = tc.get_or_add_tcPr() 
        textDirection = OxmlElement('w:textDirection') 
        textDirection.set(qn('w:val'), 'btLr') # Поворот текста на 90 градусов 
        tcPr.append(textDirection)

def getRating(id, ratings, part):
    for i in ratings:
      if(id==i['journal']['teachers']['id'] and part==i['part'] and i['rating']!=-1):
        return i
      
    return {'rating': '-'}

def getTotal(ratings):
    if(0!=len(ratings)):
      count = 0
      for i in ratings:
        if i['rating']==-1:
               count += 5
        count += i['rating']
      
      return round(count/len(ratings),1)
    
    return '-'

def mergeRecY(table, x, y):
    if y==0:
        return table.cell(x, y)
    return table.cell(x, y).merge(mergeRecY(table, x, y-1))

months = [
        "Ýanwar",
        "Fewral",
        "Mart",
        "Aprel",
        "Maý",
        "Iýun",
        "Iýul",
        "Awgust",
        "Sentýabr",
        "Oktýabr",
        "Noýabr",
        "Dekabr"
]


app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Разрешить все источники
    allow_credentials=True,
    allow_methods=["*"],  # Разрешить все методы
    allow_headers=["*"],  # Разрешить все заголовки
)

@app.post("/save-files")
async def save_files(body = Body()):
    url = requests.get('http://localhost:3000/ratings/find/poor')
    if(url.json()["success"]):
        if(not os.path.isdir(os.getenv('USERPROFILE')+'\\Desktop\\ikilikçeler')):
                os.mkdir(os.getenv('USERPROFILE')+'\\Desktop\\ikilikçeler')
                 
        for i in url.json()["datas"]:

                tm_zone = time.gmtime()

                obj = i

                doc = DocxTemplate("sorag.docx")

                rat = requests.get(f'http://localhost:3000/enumerates/israting/{obj["id"]}')

                print(obj["journal"]["subjects"]["teacher"])

                context = {
                        "month": months[int(body["month"])],
                        "year": body["year"],
                        "count": str(int(rat.json()['datas']['key']) + 1) if not rat.json()['success'] else str(int(rat.json()['datas']['key'])),
                        "subject": obj["journal"]["subjects"]["name"],
                        "primary_teacher": obj["journal"]["subjects"]["teacher"]["name"] if obj["journal"]["subjects"]["teacher"] else '______________',
                        "teacher": obj["journal"]["teachers"]["name"],
                        "kurs": obj["journal"]["groups"]["course"],
                        "group": obj["journal"]["groups"]["name"],
                        "student": obj["students"]["name"]
                }

                if not rat.json()['success']:
                        requests.post(f'http://localhost:3000/enumerates', json={'ratingsId': obj["id"], 'key': str(int(rat.json()['datas']['key']) + 1)})

                doc.render(context)
                nameFile = ''.join(f'{obj["students"]["name"]} - {obj["journal"]["teachers"]["name"]} - {obj["journal"]["subjects"]["name"]}.docx'.split('\n'))
                doc.save(os.getenv('USERPROFILE')+'\\Desktop\\ikilikçeler\\'+nameFile)

        return {'success': True}
    return {'success': False, 'error': 'NO_STUDENTS'}

@app.post("/save-files/check")
async def save_files(body = Body()):
        try:
                if(not os.path.isdir(os.getenv('USERPROFILE')+'\\Desktop\\ikilikçeler')):
                        os.mkdir(os.getenv('USERPROFILE')+'\\Desktop\\ikilikçeler')
                 
                for i in body["index"]:
                        url = requests.get(f'http://localhost:3000/ratings/{i}')
                        if(url.json()["success"]==True and url.json()["datas"]!=None):
                                tm_zone = time.gmtime()

                                obj = url.json()["datas"]

                                doc = DocxTemplate("sorag.docx")


                                rat = requests.get(f'http://localhost:3000/enumerates/israting/{url.json()["datas"]["id"]}')

                                context = {
                                        "month": months[int(body["month"])],
                                        "year": body["year"],
                                        "count": str(int(rat.json()['datas']['key']) + 1) if not rat.json()['success'] else str(int(rat.json()['datas']['key'])),
                                        "subject": obj["journal"]["subjects"]["name"],
                                        "primary_teacher": '______________',
                                        "teacher": obj["journal"]["teachers"]["name"],
                                        "kurs": obj["journal"]["groups"]["course"],
                                        "group": obj["journal"]["groups"]["name"],
                                        "student": obj["students"]["name"]
                                }

                                if not rat.json()['success']:
                                        requests.post(f'http://localhost:3000/enumerates', json={'ratingsId': url.json()["datas"]["id"], 'key': str(int(rat.json()['datas']['key']) + 1)})

                                doc.render(context)
                                nameFile = ''.join(f'{obj["students"]["name"]} - {obj["journal"]["teachers"]["name"]} - {obj["journal"]["subjects"]["name"]}.docx'.split('\n'))
                                doc.save(os.getenv('USERPROFILE')+'\\Desktop\\ikilikçeler\\'+nameFile)
                                
                return {'success': True}
        except:
                return {'success': False}

@app.post("/print-files")
async def print_files(body = Body()):
    
    url = requests.get('http://localhost:3000/ratings/find/poor')

    if(url.json()["success"]):
        for i in url.json()["datas"]:
                count = requests.get('http://localhost:3000/enumerates/count/keys')

                tm_zone = time.gmtime()
                obj = i

                doc = DocxTemplate("sorag.docx")

                rat = requests.get(f'http://localhost:3000/enumerates/israting/{obj["id"]}')

                context = {
                        "month": months[int(body["month"])],
                        "year": body["year"],
                        "count": str(int(rat.json()['datas']['key']) + 1) if not rat.json()['success'] else str(int(rat.json()['datas']['key'])),
                        "subject": obj["journal"]["subjects"]["name"],
                        "primary_teacher": obj["journal"]["subjects"]["teacher"]["name"] if obj["journal"]["subjects"]["teacher"] else '______________',
                        "teacher": obj["journal"]["teachers"]["name"],
                        "kurs": obj["journal"]["groups"]["course"],
                        "group": obj["journal"]["groups"]["name"],
                        "student": obj["students"]["name"]
                }

                if not rat.json()['success']:
                        requests.post(f'http://localhost:3000/enumerates', json={'ratingsId': obj["id"], 'key': str(int(rat.json()['datas']['key']) + 1)})

                doc.render(context)

                sections = doc.sections
                for section in sections:
                        section.left_margin = docx.shared.Inches(0.5)
                        section.right_margin = docx.shared.Inches(0.5)
                        section.top_margin = docx.shared.Inches(0.25)
                        section.bottom_margin = docx.shared.Inches(0.25)

                try:
                        namedocx = ''.join(tempfile.gettempdir()+'\\'+obj["students"]["name"]+obj["journal"]["teachers"]["name"]+obj["journal"]["subjects"]["name"]+'_time.docx'.split('\n'))
                        doc.save(namedocx)

                        win32api.ShellExecute(0, "printto", namedocx, '"%s"' % win32print.GetDefaultPrinter(), ".",0)
                except:
                        namedocx = ''.join(obj["students"]["name"]+obj["journal"]["teachers"]["name"]+obj["journal"]["subjects"]["name"]+'_time.docx'.split('\n'))
                        doc.save(namedocx)

                        win32api.ShellExecute(0, "printto", namedocx, '"%s"' % win32print.GetDefaultPrinter(), ".",0)

                count = requests.post('http://localhost:3000/enumerates', json={"key": str(int(count.json()["datas"])+1), "ratingsId": obj["id"]})

        return {'success': True}
    return {'success': False, 'error': 'NO_STUDENTS'}

@app.post("/print-files/check")
async def print_files(body = Body()):
        try:
                for i in body["index"]:
                        url = requests.get(f'http://localhost:3000/ratings/{i}')
                        if(url.json()["success"]==True and url.json()["datas"]!=None):
                                count = requests.get('http://localhost:3000/enumerates/count/keys')

                                tm_zone = time.gmtime()
                                obj = url.json()["datas"]

                                doc = DocxTemplate("sorag.docx")

                                rat = requests.get(f'http://localhost:3000/enumerates/israting/{obj["id"]}')

                                context = {
                                        "month": months[int(body["month"])],
                                        "year": body["year"],
                                        "count": str(int(rat.json()['datas']['key']) + 1) if not rat.json()['success'] else str(int(rat.json()['datas']['key'])),
                                        "subject": obj["journal"]["subjects"]["name"],
                                        "primary_teacher": '______________',
                                        "teacher": obj["journal"]["teachers"]["name"],
                                        "kurs": obj["journal"]["groups"]["course"],
                                        "group": obj["journal"]["groups"]["name"],
                                        "student": obj["students"]["name"]
                                }

                                if not rat.json()['success']:
                                        requests.post(f'http://localhost:3000/enumerates', json={'ratingsId': obj["id"], 'key': str(int(rat.json()['datas']['key']) + 1)})

                                doc.render(context)

                                sections = doc.sections
                                for section in sections:
                                        section.left_margin = docx.shared.Inches(0.5)
                                        section.right_margin = docx.shared.Inches(0.5)
                                        section.top_margin = docx.shared.Inches(0.25)
                                        section.bottom_margin = docx.shared.Inches(0.25)

                                try:
                                        namedocx = ''.join(tempfile.gettempdir()+'\\'+obj["students"]["name"]+obj["journal"]["teachers"]["name"]+obj["journal"]["subjects"]["name"]+'_time.docx'.split('\n'))
                                        doc.save(namedocx)

                                        win32api.ShellExecute(0, "printto", namedocx, '"%s"' % win32print.GetDefaultPrinter(), ".",0)
                                except:
                                        namedocx = ''.join(obj["students"]["name"]+obj["journal"]["teachers"]["name"]+obj["journal"]["subjects"]["name"]+'_time.docx'.split('\n'))
                                        doc.save(namedocx)

                                        win32api.ShellExecute(0, "printto", namedocx, '"%s"' % win32print.GetDefaultPrinter(), ".",0)

                                count = requests.post('http://localhost:3000/enumerates', json={"key": str(int(count.json()["datas"])+1), "ratingsId": obj["id"]})

                return {'success': True}
        except:
                return {'success': False}
        

@app.get("/save-journal")
async def print_files():
        if(not os.path.isdir(os.getenv('USERPROFILE')+'\\Desktop\\journals')):
                os.mkdir(os.getenv('USERPROFILE')+'\\Desktop\\journals')
        groups_urls = requests.get('http://localhost:3000/groups')
        if(groups_urls.json()['success']):

                for url2 in groups_urls.json()['datas']:

                        url = requests.get(f'http://localhost:3000/journal/find/groups/{url2["id"]}')

                        if(url.json()['success'] and url.json()['datas']!=None and url.json()['datas']['header_journal']!=None and len(url.json()['datas']['header_journal'])!=0):
                                doc = Document()

                                students = [*[[ind+1, i['name'], *[[getRating(j['teachers']['id'], i['ratings'], 1)['rating'], getRating(j['teachers']['id'], i['ratings'], 2)['rating'], getRating(j['teachers']['id'], i['ratings'], 3)['rating']] for j in url.json()["datas"]["header_journal"]], getTotal(i['ratings'])] for ind, i in enumerate(url.json()["datas"]["journal"])]]

                                students_datas = []

                                for i in range(len(students)):
                                        ln = students[i][:2]
                                        for j in range(2, len(students[i])-1):
                                                ln.extend(students[i][j])
                                        ln.append(students[i][len(students[i])-1])
                                        students_datas.append(ln)

                                datas = [
                                        [f'2024-2025-aji okuw ýylynyň I ýarym ýyllygynyň II aralyk jemleme aýyndaky Maglumat tehnologiýalary we programmirleme fakultetiniň {url2["course"]} {url2["name"]} toparynyň talyplarynyň aralyk jemlemeleriniň netijeleri barada\nMAGLUMAT'],
                                        ['T/b', 'F.A.Aa.', *[i['subjects']['name'] for i in url.json()["datas"]["header_journal"]], 'Ortaça baha'],
                                        *students_datas,
                                        ['Mugallymy', None, *[i['teachers']['name'] for i in url.json()["datas"]["header_journal"]]]
                                ]

                                section = doc.sections[-1]
                                section.orientation = WD_ORIENT.LANDSCAPE

                                new_width, new_height = section.page_height, section.page_width
                                section.page_width = new_width
                                section.page_height = new_height
                                sections = doc.sections
                                for section in sections:
                                        section.left_margin = docx.shared.Inches(0.5)
                                        section.right_margin = docx.shared.Inches(0.5)
                                        section.top_margin = docx.shared.Inches(0.25)
                                        section.bottom_margin = docx.shared.Inches(0.25)

                                table = doc.add_table(rows=len(datas), cols=len(datas[3]))
                                table.style = 'Table Grid'

                                for x_ind, x in enumerate(datas):
                                        row = table.rows[x_ind]
                                        for y_ind, y in enumerate(x):
                                                cell = row.cells[y_ind]
                                                if x_ind==1:
                                                        continue
                                                if y!=None:
                                                        if x_ind==0:
                                                                cell.text = str(y)
                                                        else:
                                                                cell.text = ''.join(str(y).split('\n'))


                                merged_cell_0  = mergeRecY(table, 0, len(datas[2])-1)

                                paragraph = merged_cell_0.paragraphs[0]
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                run = paragraph.runs[0]
                                run.bold = True
                                run.font.size = Pt(12)  

                                for i in [1, len(datas)-1]:
                                        for j in range(2, len(datas[3])-1, 3):
                                                cell_1 = table.cell(i, j)
                                                cell_2 = table.cell(i, j+1)
                                                cell_3 = table.cell(i, j+2)
                                                
                                                cell_1.merge(cell_2.merge(cell_3))

                                row = table.rows[1]
                                row.cells[0].text = datas[1][0]
                                row.cells[1].text = datas[1][1]
                                datas_1_line = datas[1][2:]

                                for y in range(len(datas_1_line)):
                                        cell = row.cells[y*3+2]
                                        if y!=None:
                                                cell.text = str(datas_1_line[y])

                                row.cells[len(row.cells)-1].text = datas[1][len(datas[1])-1]


                                row = table.rows[len(table.rows)-1]
                                datas_last_line = datas[len(datas)-1][2:]

                                for y in range(len(datas_last_line)):
                                        cell = row.cells[y*3+2]
                                        if y!=None:
                                                cell.text = str(datas_last_line[y])
                                                paragraph = cell.paragraphs[0]
                                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                                run = paragraph.runs[0]
                                                run.bold = True
                                                rotate_text(cell)


                                for i in table.rows[1:len(table.rows)-1]:
                                        for j in i.cells[:1]:
                                                j.width = Inches(0)

                                for j in table.rows[1].cells:
                                        j.width = Inches(0)
                                        paragraph = j.paragraphs[0]
                                        run = paragraph.runs[0]
                                        run.bold = True
                                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                        j.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                                row = table.rows[1]
                                row.height = Pt(80)
                                for j in table.rows[1].cells[2:]:
                                        rotate_text(j)

                                cell_1 = table.cell(len(datas)-1, 0)
                                cell_2 = table.cell(len(datas)-1, 1)
                                merged_cell  = cell_1.merge(cell_2)

                                row = table.rows[len(table.rows)-1]
                                row.height = Pt(80)
                                paragraph = merged_cell.paragraphs[0]
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                merged_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                run = paragraph.runs[0]
                                run.bold = True

                                paragraph = doc.add_paragraph("""Maglumat tehnologiýalary we programmirleme fakultetinin dekanynyň w.w.ý.ý.                         A.Annanepesow""")
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                run = paragraph.runs[0]
                                run.bold = True
                                run.font.size = Pt(12)  

                                doc.save(os.getenv('USERPROFILE')+'\\Desktop\\journals\\'+f'{url2["course"]} {url2["name"]}.docx')
        return {'success': True}    


@app.post("/save-journal/check")
async def print_files(body = Body()):
        if(not os.path.isdir(os.getenv('USERPROFILE')+'\\Desktop\\journals')):
                os.mkdir(os.getenv('USERPROFILE')+'\\Desktop\\journals')
        for i in body["index"]:
                url = requests.get(f'http://localhost:3000/journal/find/groups/{i}')
                if(url.json()['success'] and url.json()['datas']!=None and url.json()['datas']['header_journal']!=None and len(url.json()['datas']['header_journal'])!=0):

                        doc = Document()

                        url2 = requests.get(f'http://localhost:3000/groups/{i}')

                        students = [*[[ind+1, i['name'], *[[getRating(j['teachers']['id'], i['ratings'], 1)['rating'], getRating(j['teachers']['id'], i['ratings'], 2)['rating'], getRating(j['teachers']['id'], i['ratings'], 3)['rating']] for j in url.json()["datas"]["header_journal"]], getTotal(i['ratings'])] for ind, i in enumerate(url.json()["datas"]["journal"])]]

                        students_datas = []

                        for i in range(len(students)):
                                ln = students[i][:2]
                                for j in range(2, len(students[i])-1):
                                        ln.extend(students[i][j])
                                ln.append(students[i][len(students[i])-1])
                                students_datas.append(ln)

                        datas = [
                                [f'2024-2025-aji okuw ýylynyň I ýarym ýyllygynyň II aralyk jemleme aýyndaky Maglumat tehnologiýalary we programmirleme fakultetiniň {url2.json()["datas"]["course"]} {url2.json()["datas"]["name"]} toparynyň talyplarynyň aralyk jemlemeleriniň netijeleri barada\nMAGLUMAT'],
                                ['T/b', 'F.A.Aa.', *[i['subjects']['name'] for i in url.json()["datas"]["header_journal"]], 'Ortaça baha'],
                                *students_datas,
                                ['Mugallymy', None, *[i['teachers']['name'] for i in url.json()["datas"]["header_journal"]]]
                        ]

                        section = doc.sections[-1]
                        section.orientation = WD_ORIENT.LANDSCAPE

                        new_width, new_height = section.page_height, section.page_width
                        section.page_width = new_width
                        section.page_height = new_height
                        sections = doc.sections
                        for section in sections:
                                section.left_margin = docx.shared.Inches(0.5)
                                section.right_margin = docx.shared.Inches(0.5)
                                section.top_margin = docx.shared.Inches(0.25)
                                section.bottom_margin = docx.shared.Inches(0.25)

                        table = doc.add_table(rows=len(datas), cols=len(datas[3]))
                        table.style = 'Table Grid'

                        for x_ind, x in enumerate(datas):
                                row = table.rows[x_ind]
                                for y_ind, y in enumerate(x):
                                        cell = row.cells[y_ind]
                                        if x_ind==1:
                                                continue
                                        if y!=None:
                                                if x_ind==0:
                                                        cell.text = str(y)
                                                else:
                                                        cell.text = ''.join(str(y).split('\n'))


                        merged_cell_0  = mergeRecY(table, 0, len(datas[2])-1)

                        paragraph = merged_cell_0.paragraphs[0]
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = paragraph.runs[0]
                        run.bold = True
                        run.font.size = Pt(12)  

                        for i in [1, len(datas)-1]:
                                for j in range(2, len(datas[3])-1, 3):
                                        cell_1 = table.cell(i, j)
                                        cell_2 = table.cell(i, j+1)
                                        cell_3 = table.cell(i, j+2)
                                        
                                        cell_1.merge(cell_2.merge(cell_3))

                        row = table.rows[1]
                        row.cells[0].text = datas[1][0]
                        row.cells[1].text = datas[1][1]
                        datas_1_line = datas[1][2:]

                        for y in range(len(datas_1_line)):
                                cell = row.cells[y*3+2]
                                if y!=None:
                                        cell.text = str(datas_1_line[y])

                        row.cells[len(row.cells)-1].text = datas[1][len(datas[1])-1]


                        row = table.rows[len(table.rows)-1]
                        datas_last_line = datas[len(datas)-1][2:]

                        for y in range(len(datas_last_line)):
                                cell = row.cells[y*3+2]
                                if y!=None:
                                        cell.text = str(datas_last_line[y])
                                        paragraph = cell.paragraphs[0]
                                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                        run = paragraph.runs[0]
                                        run.bold = True
                                        rotate_text(cell)


                        for i in table.rows[1:len(table.rows)-1]:
                                for j in i.cells[:1]:
                                        j.width = Inches(0)

                        for j in table.rows[1].cells:
                                j.width = Inches(0)
                                paragraph = j.paragraphs[0]
                                run = paragraph.runs[0]
                                run.bold = True
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                j.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                        row = table.rows[1]
                        row.height = Pt(80)
                        for j in table.rows[1].cells[2:]:
                                rotate_text(j)

                        cell_1 = table.cell(len(datas)-1, 0)
                        cell_2 = table.cell(len(datas)-1, 1)
                        merged_cell  = cell_1.merge(cell_2)

                        row = table.rows[len(table.rows)-1]
                        row.height = Pt(80)
                        paragraph = merged_cell.paragraphs[0]
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        merged_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        run = paragraph.runs[0]
                        run.bold = True

                        paragraph = doc.add_paragraph("""Maglumat tehnologiýalary we programmirleme fakultetinin dekanynyň w.w.ý.ý.                         A.Annanepesow""")
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = paragraph.runs[0]
                        run.bold = True
                        run.font.size = Pt(12)  

                        doc.save(os.getenv('USERPROFILE')+'\\Desktop\\journals\\'+f'{url2.json()["datas"]["course"]} {url2.json()["datas"]["name"]}.docx')

        return {'success': True}     